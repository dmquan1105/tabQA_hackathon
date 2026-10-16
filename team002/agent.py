"""Table QA agent: the LLM inspects the workbook by writing pandas/openpyxl code that we execute,
then answers. Falls back to a single whole-table prompt when the agent produces no answer."""
from __future__ import annotations

import csv
import io
import logging
import re
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import httpx
import openpyxl
from openpyxl.utils import get_column_letter

log = logging.getLogger("agent")

MAX_STEPS = 8
CELL_TIMEOUT_S = 60
MAX_CELL_OUTPUT = 4000
PREVIEW_ROWS = 20
NAIVE_MAX_CHARS = 300_000


class LLM:
    """Minimal OpenAI-compatible chat client (OpenRouter) with retries and token accounting."""

    def __init__(self, base_url: str, api_key: str, model: str):
        self.url = base_url.rstrip("/") + "/chat/completions"
        self.model = model
        self.client = httpx.Client(timeout=300, headers={"Authorization": f"Bearer {api_key}"})
        self.prompt_tokens = 0
        self.completion_tokens = 0

    def chat(self, messages: list[dict], max_tokens: int = 2000) -> str:
        body = {"model": self.model, "messages": messages, "max_tokens": max_tokens, "temperature": 0,
                "reasoning": {"enabled": False}}
        last_error = None
        for attempt in range(6):
            try:
                resp = self.client.post(self.url, json=body)
                if resp.status_code in (429, 500, 502, 503, 504):
                    raise RuntimeError(f"HTTP {resp.status_code}: {resp.text[:200]}")
                data = resp.json()
                if "choices" not in data:
                    raise RuntimeError(str(data)[:300])
                usage = data.get("usage") or {}
                self.prompt_tokens += usage.get("prompt_tokens", 0) or 0
                self.completion_tokens += usage.get("completion_tokens", 0) or 0
                return data["choices"][0]["message"].get("content") or ""
            except Exception as exc:  # network errors, rate limits, malformed responses
                last_error = exc
                time.sleep(min(30, 2 ** attempt))
        raise RuntimeError(f"LLM call failed after retries: {last_error}")


# ---------------------------------------------------------------- workbook rendering

def workbook_preview(path: Path, rows: int = PREVIEW_ROWS) -> str:
    """Sheet names, sizes, merged ranges and the first rows with Excel row numbers / column letters."""
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== sheet '{ws.title}': {ws.max_row} rows x {ws.max_column} columns "
                   f"(A..{get_column_letter(max(ws.max_column, 1))})")
        merged = [str(r) for r in ws.merged_cells.ranges][:30]
        if merged:
            out.append("merged cells: " + ", ".join(merged))
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i > rows:
                out.append(f"... ({ws.max_row - rows} more rows)")
                break
            cells = [f"{get_column_letter(j)}={v!r}" for j, v in enumerate(row, start=1) if v is not None]
            out.append(f"row {i}: " + ("; ".join(cells)[:900] if cells else "(empty)"))
    return "\n".join(out)


def workbook_as_csv(path: Path, max_chars: int = NAIVE_MAX_CHARS) -> str:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        buf = io.StringIO()
        writer = csv.writer(buf)
        for row in ws.iter_rows(values_only=True):
            values = ["" if v is None else str(v) for v in row]
            while values and values[-1] == "":
                values.pop()
            writer.writerow(values)
        parts.append(f"### Sheet: {ws.title}\n{buf.getvalue()}")
    text = "\n".join(parts)
    return text if len(text) <= max_chars else text[:max_chars] + "\n...(truncated)"


def workbook_signature(path: Path, rows: int = 6) -> str:
    """Short description used to pick the right file when a query does not name one."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        lines = []
        for ws in wb.worksheets[:5]:
            lines.append(f"sheet '{ws.title}'")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= rows:
                    break
                vals = [str(v) for v in row if v is not None][:12]
                if vals:
                    lines.append(" | ".join(vals)[:300])
        return "\n".join(lines)[:1500]
    except Exception as exc:
        return f"(unreadable: {exc})"


# ---------------------------------------------------------------- answer handling

FINAL_RE = re.compile(r"FINAL ANSWER:\s*(.+)", re.S)
CODE_RE = re.compile(r"```(?:python|py)?\s*\n(.*?)```", re.S)


def clean_answer(text: str) -> str:
    lines = [line for line in text.strip().splitlines() if line.strip()]
    text = lines[0] if lines else ""
    text = re.sub(r"[*`]", "", text).strip()
    text = text.strip("\"'").strip()
    if text.endswith(".") and not re.search(r"\d\.$", text):
        text = text[:-1]
    return " ".join(text.split())


def extract_final(reply: str) -> str | None:
    match = FINAL_RE.findall(reply)
    return clean_answer(match[-1]) if match else None


ANSWER_RULES = """Answer format rules (the answer is graded by exact match):
- Output only the final value(s): no explanation, no units or currency symbols unless the question asks for them, no thousands separators.
- Copy numbers and labels exactly as they appear in the spreadsheet; if the question asks for rounding, round as asked (otherwise keep the table's precision; for computed values use 2 decimal places).
- Several values: separate them with ", " in the order the question asks.
- Yes/no questions: answer Yes or No."""

AGENT_SYSTEM = f"""You are an expert data analyst answering a question about an Excel workbook.
You have a Python tool. To run code, reply with exactly one ```python``` block; you will get its stdout back.
The variable XLSX_PATH (path of the workbook) is predefined; pandas (pd), numpy (np) and openpyxl are installed.
Every code block runs in a fresh process: reload whatever you need. Print only what you need (outputs are truncated to {MAX_CELL_OUTPUT} chars).

Work carefully:
- First inspect the real structure: multi-row/merged headers, section rows, subtotal/total rows, notes, several sheets, blank separators.
- Locate the exact row and column labels the question refers to; print the labels next to the values you use.
- Clean data before computing: strip whitespace in text cells, forward-fill labels that only appear on the first row of a group, convert numeric strings, ignore footnote markers.
- Do not double count totals together with their components.
- Verify the result with a second check before answering.

When you are confident, reply with a single line 'FINAL ANSWER: <answer>' and no code.

{ANSWER_RULES}"""

NAIVE_SYSTEM = f"""You are an expert data analyst. Answer the question using the spreadsheet given.
Think step by step, then put the answer on the last line as 'FINAL ANSWER: <answer>'.

{ANSWER_RULES}"""


def run_cell(code: str, xlsx_path: Path, workdir: str) -> str:
    script = Path(workdir) / "cell.py"
    prelude = ("import warnings; warnings.filterwarnings('ignore')\n"
               "import pandas as pd, numpy as np, openpyxl\n"
               f"XLSX_PATH = {str(xlsx_path)!r}\n")
    script.write_text(prelude + code)
    try:
        proc = subprocess.run([sys.executable, str(script)], cwd=workdir, capture_output=True,
                              text=True, timeout=CELL_TIMEOUT_S)
        output = proc.stdout + (f"\nSTDERR:\n{proc.stderr[-1500:]}" if proc.returncode else "")
    except subprocess.TimeoutExpired:
        output = f"TIMEOUT: the cell ran longer than {CELL_TIMEOUT_S}s"
    output = output.strip()
    if len(output) > MAX_CELL_OUTPUT:
        output = output[:MAX_CELL_OUTPUT] + "\n...(output truncated)"
    return output or "(no output)"


def code_agent(llm: LLM, question: str, xlsx_path: Path, deadline: float) -> str | None:
    messages = [
        {"role": "system", "content": AGENT_SYSTEM},
        {"role": "user", "content": f"Workbook: {xlsx_path.name}\n\nPreview:\n{workbook_preview(xlsx_path)}\n\n"
                                    f"Question: {question}"},
    ]
    with tempfile.TemporaryDirectory() as workdir:
        for step in range(MAX_STEPS):
            if time.time() > deadline:
                break
            reply = llm.chat(messages, max_tokens=2500)
            messages.append({"role": "assistant", "content": reply})
            code = CODE_RE.search(reply)
            final = extract_final(reply)
            if final and not code:
                return final
            if not code:
                messages.append({"role": "user", "content": "Reply with one ```python``` block, or 'FINAL ANSWER: ...'."})
                continue
            output = run_cell(code.group(1), xlsx_path, workdir)
            log.debug("step %d output: %s", step + 1, output[:300])
            messages.append({"role": "user", "content": f"Output:\n{output}"})
        messages.append({"role": "user", "content": "Stop exploring. Reply now with 'FINAL ANSWER: <answer>' only."})
        return extract_final(llm.chat(messages, max_tokens=300))


def naive_answer(llm: LLM, question: str, xlsx_path: Path) -> str | None:
    messages = [
        {"role": "system", "content": NAIVE_SYSTEM},
        {"role": "user", "content": f"Spreadsheet {xlsx_path.name} (CSV per sheet):\n\n{workbook_as_csv(xlsx_path)}\n\n"
                                    f"Question: {question}"},
    ]
    return extract_final(llm.chat(messages, max_tokens=3000))


def answer_question(llm: LLM, question: str, xlsx_path: Path, time_budget: float) -> str:
    deadline = time.time() + time_budget
    try:
        answer = code_agent(llm, question, xlsx_path, deadline)
        if answer:
            return answer
        log.info("code agent gave no answer, falling back to whole-table prompt")
    except Exception as exc:
        log.warning("code agent failed: %s", exc)
    try:
        return naive_answer(llm, question, xlsx_path) or ""
    except Exception as exc:
        log.warning("fallback failed: %s", exc)
        return ""


# ---------------------------------------------------------------- table selection

def choose_table(llm: LLM, question: str, files: list[Path], signatures: dict[Path, str]) -> Path:
    """Pick the workbook a question refers to when the query does not name a file."""
    if len(files) == 1:
        return files[0]
    catalog = "\n\n".join(f"[{i}] {p.name}\n{signatures[p]}" for i, p in enumerate(files))
    prompt = (f"Which workbook is needed to answer the question? Reply with JSON {{\"index\": <number>}} only.\n\n"
              f"Question: {question}\n\nWorkbooks:\n{catalog}")
    try:
        reply = llm.chat([{"role": "user", "content": prompt}], max_tokens=50)
        match = re.search(r"\d+", reply)
        if match and int(match.group()) < len(files):
            return files[int(match.group())]
    except Exception as exc:
        log.warning("table selection failed: %s", exc)
    # lexical fallback
    words = set(re.findall(r"[a-z0-9]{3,}", question.lower()))
    return max(files, key=lambda p: len(words & set(re.findall(r"[a-z0-9]{3,}", signatures[p].lower()))))


def resolve_table(query: dict, xlsx_dir: Path) -> Path | None:
    for key in ("file_name", "filename", "file", "table", "xlsx", "table_file"):
        name = query.get(key)
        if name:
            for candidate in (xlsx_dir / str(name), xlsx_dir / f"{name}.xlsx", xlsx_dir / Path(str(name)).name):
                if candidate.is_file():
                    return candidate
    return None
