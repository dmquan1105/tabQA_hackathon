from __future__ import annotations

import argparse
import csv
import io
import json
import re
import time
from pathlib import Path

import httpx
from openpyxl import load_workbook


SYSTEM = (
    "You are an expert data analyst. Answer the question using the spreadsheet given. "
    "Think carefully, then output only the final answer on the last line prefixed with "
    "'FINAL ANSWER:'."
)


def cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def workbook_as_csv(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow([cell_text(value) for value in row])
            parts.append(f"## Sheet: {sheet.title}\n{output.getvalue().rstrip()}")
    finally:
        workbook.close()
    return "\n\n".join(parts)


def ask_openrouter(
    client: httpx.Client,
    api_key: str,
    model: str,
    base_url: str,
    spreadsheet: str,
    file_name: str,
    question: str,
) -> str:
    messages = [
        {"role": "system", "content": SYSTEM},
        {
            "role": "user",
            "content": (
                f"Spreadsheet `{file_name}` (CSV per sheet):\n\n{spreadsheet}"
                f"\n\nQuestion: {question}"
            ),
        },
    ]
    payload = {
        "model": model,
        "messages": messages,
        "temperature": 0,
        "max_tokens": 3000,
        "reasoning": {"enabled": False},
    }
    headers = {"Authorization": f"Bearer {api_key}"}

    error: Exception | None = None
    for attempt in range(3):
        try:
            response = client.post(
                f"{base_url.rstrip('/')}/chat/completions",
                headers=headers,
                json=payload,
            )
            response.raise_for_status()
            content = response.json()["choices"][0]["message"]["content"]
            matches = re.findall(r"FINAL ANSWER:\s*(.*)", content)
            return (matches[-1] if matches else content[-300:]).strip()
        except (httpx.HTTPError, KeyError, IndexError, json.JSONDecodeError) as exc:
            error = exc
            if attempt < 2:
                time.sleep(2**attempt)
    raise RuntimeError(f"OpenRouter request failed: {error}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx-dir", type=Path, required=True)
    parser.add_argument("--queries", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--openrouter-key", required=True)
    parser.add_argument("--openrouter-model", required=True)
    parser.add_argument("--openrouter-base-url", required=True)
    args = parser.parse_args()

    data = json.loads(args.queries.read_text(encoding="utf-8-sig"))
    queries = data.get("queries", data) if isinstance(data, dict) else data

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with httpx.Client(timeout=180) as client, args.output.open(
        "w", encoding="utf-8", newline=""
    ) as output:
        writer = csv.writer(output)
        writer.writerow(["id", "answer"])
        for query in queries:
            file_name = query.get("file_name") or query.get("table")
            question = query.get("query") or query.get("question")
            spreadsheet = workbook_as_csv(args.xlsx_dir / file_name)
            answer = ask_openrouter(
                client,
                args.openrouter_key,
                args.openrouter_model,
                args.openrouter_base_url,
                spreadsheet,
                file_name,
                question,
            )
            writer.writerow([query["id"], answer])
            output.flush()


if __name__ == "__main__":
    main()
